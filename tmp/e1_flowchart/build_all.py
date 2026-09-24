from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent
S=4;W=518
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',10*S)
bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',10*S)
experiments=[('E1',False,False,False),('E2',False,True,False),('E3',True,False,False),('E4',True,True,False),('E5',False,False,True),('E6',False,True,True),('E7',True,False,True),('E8',True,True,True)]
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.4);sec.left_margin=sec.right_margin=Inches(.5)
for sn in ['Normal','Title','Caption']:
 st=doc.styles[sn];st.font.name='Times New Roman';st.font.size=Pt(10);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(3);st.paragraph_format.line_spacing=1
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
for border in doc.styles.element.xpath('.//w:pBdr'):
 border.getparent().remove(border)
def paragraph(txt,style=None):
 p=doc.add_paragraph(txt,style);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if style=='Title':
  borders=OxmlElement('w:pBdr')
  for side in ['top','bottom','left','right','between','bar']:
   el=OxmlElement('w:'+side);el.set(qn('w:val'),'nil');borders.append(el)
  p._p.get_or_add_pPr().append(borders)
 return p
for idx,(eid,q,gan,lstm) in enumerate(experiments):
 im=Image.new('RGB',(W*S,850*S),'white');d=ImageDraw.Draw(im)
 def box(cx,y,w,h,lines,rounded=False):
  bounds=((cx-w/2)*S,y*S,(cx+w/2)*S,(y+h)*S)
  if rounded:d.rounded_rectangle(bounds,radius=12*S,fill='#F1F1F1',outline='black',width=2)
  else:d.rectangle(bounds,fill='white',outline='black',width=2)
  for i,line in enumerate(lines):
   ft=bold if i==0 else font
   assert d.textlength(line,font=ft)<(w-12)*S,(eid,line)
   d.text((cx*S,(y+h/2+(i-(len(lines)-1)/2)*12)*S),line,font=ft,anchor='mm',fill='black')
 def arrow(points):
  pts=[(int(x*S),int(y*S)) for x,y in points];d.line(pts,fill='black',width=3)
  x,y=pts[-1];a,b=pts[-2]
  if y>b:tri=[(x,y),(x-3*S,y-5*S),(x+3*S,y-5*S)]
  elif x>a:tri=[(x,y),(x-5*S,y-3*S),(x-5*S,y+3*S)]
  else:tri=[(x,y),(x+5*S,y-3*S),(x+5*S,y+3*S)]
  d.polygon(tri,fill='black')
 nodes=[(18,['Start']), (30,['Load prepared PANDORA data','Comments and normalized OCEAN scores']), (30,['Separate author partitions','50 training | 10 validation | 10 test'])]
 if q:nodes.append((42,['Q-learning comment selection','Learn select-or-skip policy on training comments','Apply learned policy to each author partition']))
 else:nodes.append((30,['Baseline comment selection','Retain all available comments per author']))
 nodes.append((30,['BERT encoding','Maximum 256 tokens/comment; 768-dimensional vectors']))
 if lstm:
  nodes.append((30,['Construct ordered embedding sequences','Preserve one comment vector per timestep']))
  if gan:nodes.append((42,['GAN augmentation of training sequences only','Generate paired synthetic sequences and OCEAN scores','Combine with real training data; synthetic weight = 0.35']))
  nodes.extend([(30,['Prepare sequence batches','Pad sequences and preserve valid lengths']), (42,['Train a stacked bidirectional LSTM','2 layers; 128 hidden units/direction; 5 continuous outputs','Retain model weights using validation MAE'])])
 else:
  nodes.extend([(30,['Mean pooling','One embedding vector per author']), (30,['Feature standardization','Fit scaler on real training authors only'])])
  if gan:nodes.append((42,['GAN augmentation of training data only','Generate paired pooled vectors and OCEAN scores','Apply saved scaler; combine data; synthetic weight = 0.35']))
  nodes.append((30,['Train five ElasticNet regressors','One per OCEAN trait; alpha = 0.001; l1_ratio = 0.5']))
 nodes.extend([(30,['Predict validation OCEAN scores','Use real validation authors without augmentation']), (42,['Select one decision threshold per trait','Use validation predictions only','Maximize harmonic mean of F1 and specificity']), (30,['Freeze model'+(' and thresholds' if lstm else 's, scaler and thresholds'),'Use unchanged settings for test evaluation']), (30,['Predict continuous test OCEAN scores','Use real test authors without augmentation'])])
 y=1;last_end=None
 for height,lines in nodes:
  box(259,y,95 if lines==['Start'] else 446,height,lines,lines==['Start'])
  if last_end is not None:arrow([(259,last_end),(259,y)])
  last_end=y+height;y=last_end+8
 branch_y=y+2
 for x in [128,389]:arrow([(259,last_end),(259,branch_y-5),(x,branch_y-5),(x,branch_y)])
 box(128,branch_y,240,37,['Continuous evaluation','MAE | RMSE | R-squared | Pearson r'])
 box(389,branch_y,240,37,['Apply frozen trait thresholds','Convert scores to Low or High'])
 by=branch_y+45;arrow([(389,branch_y+37),(389,by)])
 box(389,by,240,39,['Binary evaluation','Accuracy | Precision | Recall','F1 | Specificity'])
 sy=by+48
 arrow([(128,branch_y+37),(128,sy-5),(259,sy-5),(259,sy)])
 arrow([(389,by+39),(389,sy-5),(259,sy-5),(259,sy)])
 box(259,sy,380,25,['Save models, thresholds and results'])
 ey=sy+33;arrow([(259,sy+25),(259,ey)]);box(259,ey,95,18,['End'],True)
 im=im.crop((0,0,W*S,(ey+19)*S));fn=R/(eid.lower()+'_pipeline.png');im.save(fn,dpi=(288,288))
 p=paragraph(eid+' Experiment Pipeline','Title')
 if idx:p.paragraph_format.page_break_before=True
 paragraph(('Q-learning' if q else 'Baseline')+' selection + BERT + '+('GAN + ' if gan else '')+('bidirectional LSTM' if lstm else 'ElasticNet regression'))
 p=paragraph('');p.paragraph_format.space_after=Pt(0)
 # Cap diagram height to leave title and explanatory note on the same A4 page.
 width=min(7.2,9.3*im.width/im.height)
 run=p.add_run();run.add_picture(str(fn),width=Inches(width))
 run._r.xpath('.//wp:docPr')[0].set('descr',eid+' pipeline: '+('learned Q selection' if q else 'all available comments')+', BERT, '+('training-only GAN augmentation, ' if gan else '')+('bidirectional LSTM' if lstm else 'ElasticNet')+', validation thresholds and held-out test evaluation.')
 paragraph('Figure '+str(idx+1)+': '+eid+' experimental training, validation and test pipeline.','Caption')
 p=doc.add_paragraph('Ground-truth Low/High labels use training medians. Decision thresholds use validation data. This is the experimental path; the separate saved-prediction interface embeds the first ten comments per author.')
 p.paragraph_format.space_after=Pt(0)
 for run in p.runs:run.font.size=Pt(9)
 print(eid,'height',round(im.height/im.width*width,2),'width',round(width,2))
doc.core_properties.title='Eight Experiment Pipeline Flowcharts';doc.core_properties.author=''
out=R.parents[1]/'output/documents/all_eight_experiment_flowcharts.docx';doc.save(out);print(out)
