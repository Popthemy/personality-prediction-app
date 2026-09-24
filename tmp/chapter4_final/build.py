from pathlib import Path
from copy import deepcopy
import json,hashlib,statistics
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;W=R.parents[1]
D=json.loads((R/'analysis.json').read_text());M=D['models'];P=D['predictions'];T=['Openness','Conscientiousness','Extraversion','Agreeableness','Neuroticism']
source=Path('C:/Users/STADIUM B C/Downloads/rewritten_chapter_one.docx')
doc=Document(source)
for el in list(doc._element.body):
 if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)
for s in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[s];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
 if st.element.pPr is not None:
  for el in list(st.element.pPr):
   if el.tag==qn('w:pBdr'):st.element.pPr.remove(el)
st=doc.styles['Normal'];st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_after=Pt(5);st.paragraph_format.first_line_indent=Inches(.5)
for s in ['Heading 1','Heading 2','Heading 3']:
 st=doc.styles[s];st.font.bold=True;st.paragraph_format.first_line_indent=Inches(0);st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
def p(text,style=None):
 x=doc.add_paragraph(text,style);x.paragraph_format.widow_control=True
 if not style:x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return x
def h(text,level=1):return p(text,'Heading '+str(level))
def note(text):
 x=p(text);x.paragraph_format.first_line_indent=Inches(0);x.paragraph_format.line_spacing=1.1
 for r in x.runs:r.font.size=Pt(10)
 return x
table_no=0
def table(title,heads,rows,widths=None):
 global table_no
 table_no+=1
 x=p(f'Table 4.{table_no}: {title}','Caption');x.paragraph_format.keep_with_next=True;x.paragraph_format.first_line_indent=Inches(0);x.paragraph_format.space_before=Pt(7)
 tb=doc.add_table(rows=1,cols=len(heads));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
 width=6.25
 if widths is None:widths=[width/len(heads)]*len(heads)
 for c,w in zip(tb.columns,widths):c.width=Inches(w)
 for c,t,w in zip(tb.rows[0].cells,heads,widths):c.text=t;c.width=Inches(w)
 rep=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(rep)
 for row in rows:
  cells=tb.add_row().cells
  for c,t,w in zip(cells,row,widths):c.text=str(t);c.width=Inches(w)
 for ri,row in enumerate(tb.rows):
  trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  for ci,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   cp.append(b);mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'60');el.set(qn('w:type'),'dxa');mar.append(el)
   cp.append(mar)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');cp.append(shade)
   for para in c.paragraphs:
    if ri==0:para.paragraph_format.keep_with_next=True
    para.paragraph_format.first_line_indent=Inches(0);para.paragraph_format.line_spacing=1.05;para.paragraph_format.space_after=Pt(1.5);para.paragraph_format.space_before=Pt(1.5)
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
    for run in para.runs:run.font.size=Pt(10);run.font.name='Times New Roman';run.bold=ri==0
 note('')
 return table_no
def f(x,n=4):return 'NR' if x is None else f'{x:.{n}f}'
def mean(rows,k):return sum(x[k] for x in rows)/len(rows)
latest={}
for a in P:latest.setdefault(a['condition'],a)
title=p('CHAPTER FOUR','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title=p('SYSTEM IMPLEMENTATION AND RESULTS','Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
h('4.0 Introduction')
p('This chapter explains how the design presented in Chapter Three was implemented as a working personality prediction application and evaluates the resulting eight model pipelines. Python, Django, SQLite, BERT, scikit-learn and PyTorch support the web interface, data storage and machine learning workflow. The chapter describes the development environment, interacting subsystems, dataset and user interfaces, and compares a saved training session with the latest eight successful prediction batches. Continuous OCEAN prediction is the primary evaluation task; threshold-based Low and High classification provides a secondary view of performance.')
h('4.1 Choice of Development Environment and System Requirements')
p('The application uses a Python server with Django views and templates. This environment brings data processing, model execution and result presentation into one application while keeping the machine learning components in separate service modules. HTML, CSS and JavaScript provide the browser interface. The local development database is SQLite, and model weights, embeddings and detailed reports are retained in separate artifact and cache folders. The local research workflow therefore does not require a separate database server.')
table('Software components and their implementation roles',['Component','Specification and role'],[
['Programming languages','Python 3.11 in the Dockerfile for server and models; HTML, CSS and JavaScript for interfaces'],['Web framework','Django 5.0.6 in requirements.txt; server-rendered templates and authentication'],['Database','SQLite for the local db.sqlite3; relational experiment and prediction records'],['Representation','Hugging Face Transformers; bert-base-uncased; 768 features per comment'],['Numerical libraries','NumPy, pandas and SciPy for arrays, tabular input and numerical processing'],['Learning libraries','scikit-learn for ElasticNet; PyTorch for LSTM and GAN'],['Declared ML versions','torch 2.1.2, transformers 4.35.2, scikit-learn 1.3.2 and pandas 2.1.4 in requirements-ml.txt'],['Development editor','Visual Studio Code or another Python-capable IDE can edit the project; the training log does not identify an IDE'],['Background execution','Python thread for the PANDORA experiment; Celery and Redis support other project workflows'],['Client software','A current browser with JavaScript enabled; Microsoft Word or compatible reader for exported reports']],[1.45,4.8])
p('These package versions describe dependency declarations rather than a verified inventory of the training machine. The general requirements file pins NumPy 2.4.5, whereas the machine learning requirements file pins NumPy 1.26.3. They cannot both describe one installed version. Reproduction requires a tested, consistent environment and a saved package list. The experiment artifacts do not preserve the full installed versions, processor specification or GPU model.')
table('Practical starting specifications for local operation',['Resource','Viewing results','Local model execution'],[
['Processor','64-bit dual-core processor','64-bit processor with at least four cores'],['Memory','4 GB RAM for the browser','16 GB RAM as a starting allocation; long sequences may require more'],['Storage','Space for downloads','At least 20 GB free for code, model downloads and a small working cache; data growth requires more'],['Operating system','Supported desktop OS','64-bit Windows or Linux with a compatible Python environment'],['GPU','Not required','Optional compatible CUDA GPU; CPU execution is possible but may be slow'],['Network','Access to the application','Required initially for packages and pretrained BERT downloads']],[1.0,2.05,3.2])
p('Table 4.2 gives planning specifications, not experimentally certified minimum hardware. Resource needs depend on the number of authors, comment lengths, cached embeddings and the selected pipeline. In particular, the saved augmented baseline LSTM took more than fifteen hours in this session. A minimum configuration for the largest workload can only be established by measuring memory and runtime on that workload. The browser requirements should therefore be distinguished from the server resources needed to train the models.')
h('4.2 Implementation Architecture and Subsystem Modules')
h('4.2.1 Client and server interaction',2)
p('The system follows a client-server architecture implemented with Django\'s model-template-view organization. A browser sends an authenticated request to a URL route. The corresponding view validates the form, reads or creates database records and calls the required processing service. The template renders the resulting context as a page. Django models provide persistent records; the trained prediction models are separate machine learning objects loaded from saved artifacts.')
p('A training request creates a PANDORA experiment record and starts a background Python thread. The experiment runner prepares the author partitions, selects comments, creates BERT representations and fits the eight conditions. When processing ends, the view stores condition results and artifact locations for later retrieval. This arrangement allows the researcher to return to a status page while processing continues, but the thread still depends on the application process remaining alive.')
table('Subsystems and information passed between them',['Subsystem','Input and responsibility','Output or next component'],[
['Accounts and forms','Credentials and experiment or prediction settings','Authenticated, validated request'],['Dataset loader','Three prepared PANDORA workbooks','Author IDs, comment collections and five labels'],['Comment selector','Cleaned author comments; baseline or Q-learning policy','Selected comments for BERT'],['BERT encoder','Selected text, maximum 256 tokens per comment','768-dimensional comment vectors'],['Sparse branch','Mean vector for each author','ElasticNet estimates of five continuous traits'],['Sequence branch','Ordered comment-vector sequences','Two-layer bidirectional LSTM estimates'],['GAN augmenter','Training embeddings and training scores only','Synthetic paired data for augmented fitting'],['Metrics and thresholds','True labels, scores and validation thresholds','Regression and binary classification measures'],['Persistence and interface','Runs, conditions, profiles and artifact paths','Experiment history, result tables and profile pages']],[1.25,2.7,2.3])
p('The baseline training selector retains the available comment history. The Q-learning selector learns select-or-skip decisions from training comments using text informativeness, redundancy and a selection cost; its reward is not the final personality error. Selected text then reaches BERT. The sparse branch averages the comment vectors, while the LSTM branch retains a sequence. GAN augmentation is applied to training representations and labels only. It does not add independent validation or test authors.')
p('The saved condition names use the prefix lasso, but the selected session explicitly configures ElasticNet with alpha 0.001 and l1_ratio 0.5. The chapter therefore calls this branch sparse regression or ElasticNet. The LSTM has two bidirectional recurrent layers, 128 hidden units per direction, dropout 0.2, a learning rate of 0.001, batch size four and a maximum of 35 epochs. The GAN configuration specifies 150 epochs, latent dimension 64, hidden dimension 128, batch size 16 and learning rate 0.0002. Each augmented condition records 50 generated rows. These are derived training examples, not additional people.')
h('4.2.2 Dataset source and experimental allocation',2)
p('The source is PANDORA, a Reddit comment dataset linked with personality and demographic labels. Gjurković et al. (2021) describe approximately 10,000 users overall, including approximately 1,600 with Big Five labels. These publication-level totals describe the source corpus, not the number used in this project\'s selected experiment. The immediate project inputs are Training dataset.xlsx, Validation dataset.xlsx and Test dataset.xlsx in the PANDORA folder.')
p('The selected session is database experiment 17, pandora_20260916_202608, with artifacts in run_20260917_133820. Its 50 training authors satisfy the requested upper limit of 200 training users. All eight conditions share 50 training, ten validation and ten test authors, with random seed 42. File-defined partitions determine membership; the actual allocation is 71.43%, 14.29% and 14.29% of the 70 sampled authors. It must not be described as a newly applied 60:20:20 split merely because ratio settings remain in the configuration.')
rows=[]
for k,label in [('train','Training'),('val','Validation'),('test','Test')]:
 q=D['quality']['comment_volume'][k];rows.append([label,q['n_participants'],q['n_comments'],f(q['mean_comments_per_participant'],1),f(q['median_comments_per_participant'],1),f"{q['min_comments']}–{q['max_comments']}"])
rows.append(['Total',70,'11,019','157.4','—','—'])
table('Authors and comments in the selected training session',['Split','Authors','Comments','Mean per author','Median','Range'],rows,[1.1,.65,.9,1.2,.95,1.45])
p('The experiment contains 11,019 comments: 8,675 in training, 681 in validation and 1,663 in test. The saved overlap checks report zero common authors for every pair of partitions. Authors, rather than comments, are the independent labelled observations. Comment histories vary greatly, so the comment total does not imply 11,019 independent personality labels. The runner loaded the files as already cleaned and reports no additional comment removals during this run; this does not describe all earlier preparation of the source dataset.')
p('The source scores are divided by 100 to obtain values on the zero-to-one scale. Binary ground truth uses training medians: 0.7900 for Openness, 0.3550 for Conscientiousness, 0.2500 for Extraversion, 0.3500 for Agreeableness and 0.4600 for Neuroticism. A score at or above its boundary is High. These ground-truth boundaries are different from the model-specific decision thresholds selected on validation data.')
h('4.2.3 Storage and the saved prediction path',2)
p('The principal database entities are experiment runs, condition results, threshold results, dataset allocations, prediction runs and test profiles. A prediction run links to its source experiment and selected condition. Each test profile retains the observed and predicted continuous traits, binary labels, decision thresholds and correctness flags. Consequently, the aggregate prediction measures can be recalculated from the stored profiles rather than accepted only from a dashboard summary.')
p('The operational prediction helper embeds the first ten non-empty cleaned comments for every author. It does not reapply the training experiment\'s full-history baseline or learned Q-selection policy. Pipeline names in prediction records therefore identify the fitted model\'s training condition; they do not demonstrate identical comment selection at inference. This difference is material when interpreting a training-to-prediction change or attributing a gain specifically to Q-learning.')
h('4.3 System Interface and Screen Shots')
p('The interface separates authentication, experiment execution and prediction inspection. The following descriptions identify the principal implemented pages and the content required in their corresponding figures. Performance conclusions are based on saved research records rather than interface appearance.')
h('4.3.1 Login and researcher dashboard',2)
p('The login page controls access to the research functions through Django authentication. After authentication, the dashboard summarizes available records and recent PANDORA runs. It provides navigation to the tools used for training and prediction. Personal account details should be hidden in any reproduced interface image.')
note('Figure 4.1: Login page and researcher dashboard — screenshot pending.')
h('4.3.2 Experiment setup and results',2)
p('The training interface accepts experiment settings and starts a PANDORA run. The experiment history lists saved runs, while the detail page identifies the run, its state and the eight condition results. A suitable screenshot should show experiment 17 and its completed results so that the interface evidence corresponds to the numerical tables in this chapter.')
note('Figure 4.2: Training setup and completed eight-condition experiment — screenshot pending.')
h('4.3.3 Prediction results and individual profiles',2)
p('The prediction results page displays the batch identifier, fitted condition, number of profiles and summary classification metrics. Its displayed average decision threshold is only a summary of five thresholds; prediction uses a separate threshold for each trait. The profile page compares observed and predicted OCEAN values and their Low or High labels. These views support inspection of individual errors as well as overall scores.')
note('Figure 4.3: Saved prediction batch and anonymized OCEAN profile — screenshot pending.')
h('4.4 Experiments and Results')
h('4.4.1 Experimental conditions and evidence scope',2)
p('The experiment varies prediction model, selection method and augmentation in a two-by-two-by-two design. E1–E8 below are short labels used consistently in the tables. They represent eight fitted conditions, not eight folds of cross-validation. The selected source experiment contains 70 real authors in total; each model is fitted with the same 50 real training authors.')
table('Definition of the eight experiments',['ID','Saved condition','Predictor','Selection','GAN'],[[v['id'],k,'ElasticNet' if k.startswith('lasso') else 'BiLSTM','Q-learning' if 'qlearn' in k else 'Baseline','Yes' if k.endswith('gan') else 'No'] for k,v in M.items()],[.4,2.15,1.2,1.0,1.5])
p('The prediction comparison uses the latest eight successful prediction batches, ordered by creation time and record ID. The failed record 36 is excluded. The resulting records are 40, 39, 38, 37, 35, 34, 33 and 32. All link to experiment 17. They contain seven distinct pipelines because E6 occurs twice; E2 has no result within this requested window. E2 is retained in the training tables and marked unavailable in prediction comparisons rather than assigned a result from an older batch.')
p('Seven batches contain the same 33 authors. Record 33 contains a 20-author subset. These batches reuse profiles and are not eight independent test samples. The 33-author set has no overlap with the training or validation authors, but includes all ten original test authors. It therefore extends the observed test sample rather than providing a completely independent external test. Comparisons used to choose a pipeline are restricted to the equal 33-author batches.')
h('4.4.2 Performance measures and reporting conventions',2)
p('Mean absolute error (MAE) is the mean absolute difference between the observed and predicted continuous scores. Root mean squared error (RMSE) gives more weight to large errors. Lower MAE and RMSE indicate better predictions. R² compares squared error with prediction by the evaluated sample mean; a negative value indicates worse performance than that reference. Pearson correlation summarizes linear association and is reported as zero for constant predictions to match the project convention.')
p('For binary evaluation, TP and TN denote correctly predicted High and Low labels, while FP and FN denote false High and false Low labels. Accuracy = (TP + TN)/N; precision = TP/(TP + FP); recall = TP/(TP + FN); specificity = TN/(TN + FP); and High-class F1 = 2TP/(2TP + FP + FN). Undefined precision or F1 denominators are assigned zero, consistent with the saved prediction function. ROC AUC and PR AUC (average precision in the saved implementation) describe score ranking across operating points rather than performance at only the selected threshold.')
p('Validation tables use an unweighted mean across the five traits. The application\'s saved prediction summary pools all author-trait binary decisions before calculating precision, recall, F1 and specificity. These pooled quantities differ from the mean of five trait metrics. The chapter presents both forms with explicit labels and uses trait means for training-to-prediction comparisons. RMSE is likewise averaged over the five trait-specific RMSE values; it is not a single square root after pooling all trait errors.')
p('Some saved sparse-regression accuracy and macro-F1 fields refer to three classes, whereas the threshold fields refer to binary Low and High classes. Copying them into one binary comparison would mix definitions. The binary validation tables below use the retained binary precision, recall, F1 and specificity. Sparse binary accuracy is recovered from the saved positive and negative class counts, recall and specificity; the resulting confusion counts are integers. LSTM binary values come from the validation sweep row at the chosen threshold. The test-stage val_mae field is identified by its containing test block and is reported as test MAE, not validation MAE.')
h('4.4.3 Model fitting and overall validation performance',2)
rows=[]
for k,m in M.items():
 vals=[r['validation'].get('train_mae') for r in m['rows']];train=statistics.mean(vals) if all(x is not None for x in vals) else None
 rows.append([m['id'],50,50 if k.endswith('gan') else 0,f(train),f(m['validation']['mae']),f(m['test']['mae']),f(m['seconds'],2)])
table('Stored training-session performance and fitting cost',['ID','Real train authors','Synthetic rows','Train MAE','Val MAE','Test MAE','Fit seconds'],rows,[.4,.8,.8,1.0,1.0,1.0,1.25])
note('NR = not recorded. Training MAE is the stored branch diagnostic; augmented diagnostics are not asserted to be measured on real authors alone. The source does not retain equivalent training-set metrics or epoch histories for the LSTM conditions. Validation and test values are explicitly separated from in-sample training performance.')
p('The lowest validation MAE is E5 at 0.24698, while the lowest original ten-author test MAE is E3 at 0.21934. This difference shows that the preferred condition depends on the evaluated authors. The absence of saved LSTM training-set histories prevents a complete learning-curve or overfitting comparison. Training time is recorded by condition and is descriptive of this run; cache state and machine load were not controlled sufficiently to interpret it as a general hardware benchmark.')
table('All pipelines across validation regression and ranking metrics',['ID','MAE','RMSE','R²','Pearson r','ROC AUC','PR AUC'],[[m['id']]+[f(m['validation'].get(k)) for k in ['mae','rmse','r2','val_pearson','roc_auc','pr_auc']] for m in M.values()],[.4,.9,.95,.9,1.1,1.0,1.0])
table('All pipelines across consistent binary validation metrics',['ID','Accuracy','Precision','Recall','F1 High','Specificity'],[[m['id']]+[f(mean([r['binary_validation'] for r in m['rows']],k)) for k in ['accuracy','precision','recall','f1','specificity']] for m in M.values()],[.4,1.1,1.15,1.1,1.2,1.3])
p('Each validation trait is represented by only ten authors, so moving one classification changes its accuracy by 0.10. Threshold tuning and reporting use the same validation authors. Validation classification scores are consequently development results, not unbiased estimates from an untouched sample. The separate ten-author test and later saved predictions must be considered alongside them.')
h('4.4.4 Validation performance across the five traits',2)
rows=[]
for m in M.values():
 for r in m['rows']:
  v=r['validation'];rows.append([m['id'],r['trait'][0]]+[f(v.get(k)) for k in ['train_mae','val_mae','val_rmse','val_r2','val_pearson']])
table('Per-trait training diagnostics and validation regression performance',['ID','Trait','Train MAE','Val MAE','Val RMSE','Val R²','Val r'],rows,[.4,.5,1.05,1.05,1.1,1.1,1.05])
note('O = Openness; C = Conscientiousness; E = Extraversion; A = Agreeableness; N = Neuroticism. All validation rows use ten authors. NR indicates an absent stored training metric, not zero error.')
p('The trait breakdown prevents an overall mean from hiding uneven performance. For example, E1 has validation MAE 0.2130 for Openness and 0.2919 for Neuroticism. Its constant predictions on several other traits have Pearson correlation recorded as zero. Such results indicate that a pipeline may learn some variation in one target while effectively predicting a common score for another. Claims about the entire OCEAN profile therefore require all five traits, not only the best-performing trait.')
h('4.4.5 Selection of a decision threshold for each trait',2)
p('Threshold selection begins after continuous prediction. The training-median boundary creates the observed High or Low label, while a candidate decision threshold converts the model\'s validation score into a predicted label. With ten validation scores per trait, the implementation uses the distinct predicted score values, clipped to the zero-to-one interval and rounded to four decimal places, as candidates. The general implementation uses percentiles when there are more than nineteen scores. The fixed fallback list in the configuration is therefore not the complete set of thresholds actually considered in this session.')
p('For each candidate, the selector computes binary metrics and maximizes the harmonic mean of High-class F1 and specificity: H = 2 × F1 × specificity / (F1 + specificity). H is zero when its denominator is zero. The intent is to discourage selecting High for every author simply to obtain a large recall. Ties are resolved by higher F1, then higher specificity, then higher accuracy; a remaining exact tie retains the first candidate in the sorted list. The selected validation threshold is frozen for test evaluation and loaded by the prediction function.')
rows=[]
for m in M.values():
 for r in m['rows']:
  b=r['binary_validation'];rows.append([m['id'],r['trait'][0],f(r['threshold'])]+[f(b[k]) for k in ['accuracy','precision','recall','f1','specificity']])
table('Selected trait thresholds and binary validation performance at each threshold',['ID','Trait','Threshold','Accuracy','Precision','Recall','F1','Specificity'],rows,[.4,.45,.95,.9,.9,.85,.85,.95])
p('A selected threshold is not proof of useful separation. E1, for example, records zero High-class F1 for Extraversion and Agreeableness despite specificity of 1.0000. Those decisions classify every validation example as Low. The threshold is an operational choice from the observed scores; it is not a population personality boundary or a calibrated probability cutoff. Test-only threshold sweeps are descriptive and are not used to replace the validation-selected values.')
h('4.4.6 The latest eight successful prediction batches',2)
rows=[]
for a in P:
 rows.append([a['id'],M[a['condition']]['id'],a['created_at'][11:19],a['requested_samples'],a['predicted_samples'],a['comments'],f(a['means']['mae'])])
table('Prediction-batch inventory and continuous performance',['Record','ID','Time UTC','Requested','Predicted','Source comments','MAE'],rows,[.65,.4,1.0,.85,.85,1.35,1.05])
note('All batches were created on 17 September 2026 and refer to experiment 17. Times are the stored UTC timestamps. Source comments count available comments attached to profiles, not comments embedded: the operational helper processes at most ten per author. Requested sample size does not equal the number available for prediction.')
table('Each saved prediction batch across the application classification metrics',['Record / ID','Accuracy','Precision','Recall','F1 High','Specificity'],[[str(a['id'])+' / '+M[a['condition']]['id']]+[f(a[k]) for k in ['accuracy','precision','recall','f1_score','specificity']] for a in P],[1.0,1.05,1.05,1.0,1.05,1.2])
note('These values are independently checked against the stored binary profile labels. Metrics pool five decisions per author: 165 decisions in each 33-author batch and 100 in the 20-author batch. They are not eight repetitions of a common independent experiment.')
table('Each saved prediction batch across continuous and ranking metrics',['Record / ID','n','MAE','RMSE','R²','r','ROC AUC','PR AUC'],[[str(a['id'])+' / '+M[a['condition']]['id'],a['predicted_samples']]+[f(a['means'][k]) for k in ['mae','rmse','r2','pearson','roc_auc','pr_auc']] for a in P],[.8,.55,.8,.8,.8,.8,.85,.85])
p('Prediction ROC AUC and average precision are recalculated from the saved continuous scores and binary ground truth, with ties handled consistently. The equal-size comparison gives E8 the lowest prediction MAE and RMSE, while E3 gives the highest pooled binary accuracy. E7 gives the highest pooled F1 and recall, but its specificity is only 0.1379, showing that it often predicts High for truly Low cases. Record 33 has a smaller MAE than the 33-author batches, but this cannot establish E6 as the best pipeline because it uses a different 20-author subset. Both E6 records remain visible to preserve the requested latest-eight selection.')
h('4.4.7 Training and validation compared with actual prediction',2)
rows=[]
for code,m in M.items():
 rows.append([m['id'],'Validation']+[f(m['validation'].get(k)) for k in ['mae','rmse','r2','val_pearson','roc_auc','pr_auc']])
 a=latest.get(code);rows.append([m['id'],'Prediction '+str(a['id']) if a else 'Not in window']+[f(a['means'][k]) if a else '—' for k in ['mae','rmse','r2','pearson','roc_auc','pr_auc']])
table('Direct validation and prediction comparison using trait-mean regression metrics',['ID','Stage','MAE','RMSE','R²','r','ROC AUC','PR AUC'],rows,[.4,1.15,.75,.8,.8,.8,.8,.8])
rows=[]
for code,m in M.items():
 rows.append([m['id'],'Validation']+[f(mean([r['binary_validation'] for r in m['rows']],k)) for k in ['accuracy','precision','recall','f1','specificity']])
 a=latest.get(code);rows.append([m['id'],'Prediction '+str(a['id']) if a else 'Not in window']+[f(a['means'][k]) if a else '—' for k in ['accuracy','precision','recall','f1','specificity']])
table('Direct validation and prediction comparison using trait-mean binary metrics',['ID','Stage','Accuracy','Precision','Recall','F1','Specificity'],rows,[.4,1.25,.95,.95,.85,.9,1.0])
note('Prediction rows use the newest batch for each pipeline inside the latest-eight window, all with 33 authors. Record 33 is still reported in Tables 4.11–4.13 but is not substituted for the larger E6 comparison. E2 is unavailable in this window. Binary prediction F1 here is a mean over traits and differs from the pooled F1 in Table 4.12.')
p('These paired tables align metric definitions, but they do not make the evaluated populations or comment-selection procedures identical. Validation uses ten authors and the experimental feature path. Prediction uses 33 authors and the first-ten-comment helper. Changes therefore combine sample composition, threshold behaviour and an implementation difference. They cannot be attributed solely to overfitting, Q-learning or augmentation.')
h('4.4.8 Detailed outcomes of the eight experiments',2)
explain={
'lasso_baseline':'E1 is the pooled-BERT sparse baseline. It establishes the effect of fitting ElasticNet without learned selection or augmentation. Several targets have no active sparse features, which limits their ability to follow author-level variation.',
'lasso_baseline_gan':'E2 adds training-only GAN augmentation to E1. Its original test MAE is lower than E1, but its validation MAE is higher. This mixed result does not support a general claim that augmentation improves sparse prediction. There is no E2 batch in the latest eight successful predictions, so its operational performance is not ranked.',
'lasso_qlearn':'E3 uses Q-learning-selected BERT comments with ElasticNet and no GAN. It has the lowest original ten-author test MAE. In actual prediction it achieves the best pooled accuracy among the equally sized batches, but its low recall shows that this accuracy does not imply equally good identification of High labels.',
'lasso_qlearn_gan':'E4 combines sparse regression, Q-learning and GAN augmentation. Its validation MAE is better than E3, but its original test error and saved prediction error are worse. The combined additions therefore do not provide a consistent improvement over the simpler sparse condition.',
'lstm_baseline':'E5 replaces pooled regression with a bidirectional sequence model and uses the baseline comment history during training. It gives the lowest validation MAE. Its prediction error is close to E8, while its pooled F1 and recall are lower and its specificity is higher.',
'lstm_baseline_gan':'E6 adds GAN augmentation to the baseline LSTM. Its original test RMSE is slightly lower than the other conditions, but its recorded fitting duration is exceptionally long. The two saved prediction batches demonstrate why results from 20 and 33 authors must not be treated as matched repetitions.',
'lstm_qlearn':'E7 uses the sequence model trained on Q-selected comments without augmentation. Its prediction F1 and recall lead the equal-size saved batches, but low specificity reveals a pronounced tendency to predict High. This is a classification trade-off rather than a general improvement in all measures.',
'lstm_qlearn_gan':'E8 combines Q-learning, BERT, GAN augmentation and the bidirectional LSTM. It gives the lowest MAE and RMSE on the common 33-author prediction set. This supports its selection for continuous OCEAN estimation within the observed comparison, while the negative mean R² and low specificity limit the strength and scope of the recommendation.'}
for code,m in M.items():
 h(m['id']+' '+('ElasticNet' if code.startswith('lasso') else 'Bidirectional LSTM')+' with '+('Q learning selection' if 'qlearn' in code else 'baseline selection')+(' and GAN' if code.endswith('gan') else ''),3)
 p(explain[code])
 val=m['validation'];te=m['test'];b={k:mean([r['binary_validation'] for r in m['rows']],k) for k in ['accuracy','precision','recall','f1','specificity']}
 p(f"Validation MAE was {f(val['mae'])}, RMSE {f(val['rmse'])}, R² {f(val['r2'])} and Pearson r {f(val['val_pearson'])}. Mean binary validation accuracy was {f(b['accuracy'])}, precision {f(b['precision'])}, recall {f(b['recall'])}, F1 {f(b['f1'])} and specificity {f(b['specificity'])}. The original test MAE was {f(te['mae'])}, compared with a recorded fitting time of {f(m['seconds'],2)} seconds. Trait-level errors and threshold-specific metrics are given in Tables 4.9 and 4.10.")
 a=latest.get(code)
 if a:
  p(f"In prediction record {a['id']} ({a['predicted_samples']} authors), MAE was {f(a['means']['mae'])}, RMSE {f(a['means']['rmse'])}, mean R² {f(a['means']['r2'])} and mean Pearson r {f(a['means']['pearson'])}. The application's pooled accuracy was {f(a['accuracy'])}, precision {f(a['precision'])}, recall {f(a['recall'])}, F1 {f(a['f1_score'])} and specificity {f(a['specificity'])}. Tables 4.14 and 4.15 give the corresponding validation-to-prediction comparison under common trait-mean definitions.")
h('4.4.9 Preferred pipeline during prediction',2)
winner=next(a for a in P if a['id']==38)
p('E8, the model trained with Q-learning selection, BERT representations, GAN augmentation and a bidirectional LSTM, is the preferred pipeline for continuous OCEAN estimation among the seven pipelines represented on the common 33-author prediction set. The selection rule prioritizes mean MAE, then mean RMSE. E8 records MAE 0.2466 and RMSE 0.2873, both the lowest in that matched comparison. This choice follows the primary continuous-score objective established in Chapters One and Three. It is not an assertion that E8 wins every classification metric or that the unrepresented E2 has been ruled out.')
e5=latest['lstm_baseline'];e6=latest['lstm_baseline_gan'];e3=latest['lasso_qlearn']
p(f"E8 improves MAE over E5 by {f(e5['means']['mae']-winner['means']['mae'],6)} and over E6 by {f(e6['means']['mae']-winner['means']['mae'],6)} on the same authors. Its improvement over E3 is {f(e3['means']['mae']-winner['means']['mae'],6)}. These differences are descriptive: the small E8–E5 margin is not evidence of statistical significance. No repeated-seed analysis or confidence interval establishes a stable population ranking. All compared mean R² values remain negative, including E8 at {f(winner['means']['r2'])}, so the selected pipeline still has weak continuous predictive performance.")
table('Selected E8 thresholds and actual per-trait prediction results',['Trait','Label cutoff','Decision threshold','MAE','Accuracy','F1','Specificity'],[[t[0],f(D['imbalance']['splits']['train']['traits'][t[0]]['binary_cutoff']),f(winner['thresholds'][t.lower()]),f(winner['traits'][i]['mae']),f(winner['traits'][i]['accuracy']),f(winner['traits'][i]['f1']),f(winner['traits'][i]['specificity'])] for i,t in enumerate(T)],[.5,1.0,1.2,.85,.9,.85,.95])
p('The operational thresholds for E8 are 0.6837 for Openness, 0.4523 for Conscientiousness, 0.3491 for Extraversion, 0.3980 for Agreeableness and 0.4696 for Neuroticism. They match the thresholds stored in its prediction profiles and originate from validation. Their arithmetic mean, 0.47054, is the value displayed at batch level; it should not replace the five thresholds in individual trait decisions.')
table('Full E8 per-trait prediction metrics at the selected operating thresholds',['Trait','RMSE','R²','Precision','Recall','TP / FP','TN / FN'],[[r['trait'][0],f(r['rmse']),f(r['r2']),f(r['precision']),f(r['recall']),str(r['tp'])+' / '+str(r['fp']),str(r['tn'])+' / '+str(r['fn'])] for r in winner['traits']],[.5,.95,.95,1.0,.95,.95,.95])
p('Across all 165 author-trait decisions, E8 produces 58 true High predictions, 19 true Low predictions, 68 false High predictions and 20 false Low predictions. Its pooled accuracy is 0.4667, F1 is 0.5686 and specificity is 0.2184. A classifier that always predicts Low on this set would obtain 87/165 = 0.5273 accuracy but zero High-class recall and F1. Consequently, the selected continuous predictor cannot also be claimed to provide strong binary classification. E3 is preferable if the narrow criterion is observed binary accuracy among these pipelines, and E7 is preferable if the criterion is observed High-class F1; neither conclusion establishes general deployment quality.')
h('4.4.10 Discussion of component effects and study limitations',2)
p('Q-learning reduced the experiment\'s selected comment total from 11,019 to 7,598, a reduction of 3,421 comments or 31.05%. The average decreased from 157.41 to 108.54 comments per author. This is evidence of lower input volume during the experimental feature path, not a measured inference-time saving in the first-ten-comment helper. The recorded training times also vary substantially between conditions and include an unusually large 56,253.64 seconds for E6.')
p('On the original test set, adding Q-learning without GAN reduced MAE for both the sparse and sequence branches, while adding it with GAN increased MAE for both branches. Augmentation likewise helped the baseline-selection test conditions but harmed the Q-learning test conditions. The directions differ again in the later prediction comparison. These interactions show why the combined pipeline must be evaluated directly instead of assuming that two potentially useful components will necessarily improve one another.')
p('The strongest limitations are the 50-author training sample, ten-author validation set, one recorded random seed and reused test authors. Thresholds optimized on ten validation scores can be sensitive to individual authors. The saved source repository was marked dirty, so its commit identifier does not fully identify all code executed. Training histories and the full hardware inventory were not retained. Most importantly, operational inference uses a different comment-selection procedure from the source experiment. These constraints limit causal attribution and reproducibility, even though the saved prediction summaries can be checked against the stored profile labels.')
p('The results support the implementation of an integrated eight-condition research system and a provisional selection of E8 for continuous prediction within the observed equal-size batches. They do not establish a universally best pipeline, reliable population-level personality assessment or a complete eight-pipeline operational comparison. Further evaluation should preserve raw predictions, reproduce each condition\'s selection path at inference, include E2 on the same held-out authors, and compare multiple seeds on a larger independent sample.')
h('References')
p('Gjurković, M., Karan, V. M., Vukojević, I., Bošnjak, M., & Šnajder, J. (2021). PANDORA talks: Personality and demographics on Reddit. Proceedings of the Ninth International Workshop on Natural Language Processing for Social Media, 138–152. https://doi.org/10.18653/v1/2021.socialnlp-1.12')
p('Project evidence: SQLite experiment record 17 and successful prediction records 40, 39, 38, 37, 35, 34, 33 and 32; run_20260917_133820/experiment_data.json, data_quality.json, imbalance_report.json, trait_label_thresholds.json, reproducibility.json, run_summary.json and the eight condition metrics.json files. Results inspected on 18 September 2026. Implementation references: backend/tools/views.py, backend/core/models.py, backend/ml_pipeline/experiments/pandora_runner.py and backend/ml_pipeline/services/metrics_engine.py.')
doc.core_properties.title='Chapter Four System Implementation and Results';doc.core_properties.author='';doc.core_properties.subject='Eight experimental conditions and latest eight successful prediction batches'
out=W/'output/documents/chapter_four_guideline_rewrite.docx';doc.save(out)
(R/'artifact.md').write_text('Reference: '+str(source)+'\nSHA256 '+hashlib.sha256(source.read_bytes()).hexdigest()+'\nChapter body replaced; reference retained unchanged. Letter portrait; inherited section geometry, footer, styles and Times New Roman 12 pt body. Tables use 10 pt for readable numerical comparisons, repeated headers, non-splitting rows and light gray borders. Exact user headings 4.0 to 4.4 retained. Actual screenshots unavailable because browser provider fails policy initialization; figure slots explicitly marked pending. Packaged renderer failed due to absent LibreOffice; use Word PDF export and page raster review.\n',encoding='utf8')
print(out);print('Tables',table_no,'Words',sum(len(x.text.split()) for x in doc.paragraphs))
