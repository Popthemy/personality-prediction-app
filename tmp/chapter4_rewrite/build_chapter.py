from pathlib import Path
from copy import deepcopy
import re, zipfile, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
SOURCE=Path('C:/Users/STADIUM B C/Downloads/rewritten_chapter_one.docx')
OUT=ROOT.parents[1]/'output/documents/chapter_four_revised.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc=Document(SOURCE)
props={k:deepcopy(doc.paragraphs[i]._p.pPr) for k,i in [('title',0),('heading',2),('body',3)]}
for el in list(doc._element.body):
    if el.tag!=qn('w:sectPr'):doc._element.body.remove(el)

def font(run,bold=False,italic=False):
    run.font.name='Times New Roman';run.font.size=Pt(12)
    run.font.color.rgb=RGBColor(0,0,0);run.bold=bold;run.italic=italic

def paragraph(text,kind='body'):
    p=doc.add_paragraph();p._p.insert(0,deepcopy(props[kind]))
    p.paragraph_format.widow_control=True
    if kind=='title':
        p.style=doc.styles['Title'];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Inches(0)
        borders=OxmlElement('w:pBdr')
        for side in ['top','left','bottom','right','between','bar']:
            el=OxmlElement('w:'+side);el.set(qn('w:val'),'nil');borders.append(el)
        p._p.pPr.append(borders)
    elif kind=='heading':
        p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.keep_with_next=True
        el=OxmlElement('w:outlineLvl');el.set(qn('w:val'),str(min(2, len(re.match(r'[0-9.]+', text).group().rstrip('.').split('.'))-2)) if re.match(r'4\.',text) else '0');p._p.pPr.append(el)
    font(p.add_run(text),bold=kind!='body',italic=text.startswith('['))
    return p

in_refs=False
for block in ROOT.joinpath('chapter_four.md').read_text(encoding='utf-8').split('\n\n'):
    block=block.strip()
    if not block:continue
    if block.startswith('# CHAPTER'):
        paragraph('CHAPTER FOUR','title')
        paragraph('SYSTEM IMPLEMENTATION, TESTING, RESULTS AND DISCUSSION','title')
    elif block.startswith('## '):
        p=paragraph(block[3:],'heading')
        if block=='## REFERENCES':
            in_refs=True;p.paragraph_format.page_break_before=True
    elif block.startswith('### '):paragraph(block[4:],'heading')
    elif block.startswith('#### '):paragraph(block[5:],'heading')
    else:
        p=paragraph(block)
        if block.startswith('['):
            p.paragraph_format.first_line_indent=Inches(0)
            p.paragraph_format.keep_together=True
        if in_refs:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent=Inches(-.5);p.paragraph_format.left_indent=Inches(.5)
            p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
            spans={'Devlin':'Proceedings of the 2019 Conference of the North American Chapter of the Association for Computational Linguistics: Human Language Technologies','Friedman':'Journal of Statistical Software, 33','Goodfellow':'Advances in Neural Information Processing Systems','Sutton':'Reinforcement learning: An introduction'}
            for prefix,span in spans.items():
                if block.startswith(prefix):
                    p.runs[0].text='';a,b=block.split(span,1)
                    font(p.add_run(a));font(p.add_run(span),italic=True);font(p.add_run(b))

doc.core_properties.title='Chapter Four: System Implementation, Testing, Results and Discussion'
doc.core_properties.subject='Personality prediction: implementation, testing, results and discussion'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
doc.save(OUT)
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','word/settings.xml','docProps/core.xml','docProps/app.xml'}
with zipfile.ZipFile(SOURCE) as z:src={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(OUT) as z:dst={n:z.read(n) for n in z.namelist()}
for n,b in src.items():
    if n not in editable and n in dst:dst[n]=b
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
    for n,b in dst.items():z.writestr(n,b)
assert all(dst[n]==b for n,b in src.items() if n not in editable and n in dst)
ROOT.joinpath('source_inventory.json').write_text(json.dumps({n:hashlib.sha256(b).hexdigest() for n,b in src.items()},indent=2))
with ROOT.joinpath('artifact.md').open('a',encoding='utf-8') as f:f.write('\nReference SHA256: '+hashlib.sha256(SOURCE.read_bytes()).hexdigest()+'\nPreserve-only package parts verified identical.\n')
print(OUT)
print('Chapter words:',len(ROOT.joinpath('chapter_four.md').read_text(encoding='utf-8').split('## REFERENCES')[0].split()))
