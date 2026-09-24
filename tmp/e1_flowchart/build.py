from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
ROOT.mkdir(parents=True,exist_ok=True)
S=4
# Coordinates are in points; the final diagram is 7.2 inches wide.
W,H=518,637
im=Image.new('RGB',(W*S,H*S),'white');d=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',10*S)
bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',10*S)
def box(cx,y,w,h,lines,rounded=False):
 bounds=((cx-w/2)*S,y*S,(cx+w/2)*S,(y+h)*S)
 if rounded:d.rounded_rectangle(bounds,radius=12*S,fill='#F1F1F1',outline='black',width=2)
 else:d.rectangle(bounds,fill='white',outline='black',width=2)
 for i,line in enumerate(lines):
  ft=bold if i==0 else font
  d.text((cx*S,(y+h/2+(i-(len(lines)-1)/2)*12)*S),line,font=ft,anchor='mm',fill='black')
 return (cx,y,w,h)
def arrow(points):
 p=[(int(x*S),int(y*S)) for x,y in points];d.line(p,fill='black',width=3)
 x,y=p[-1];x0,y0=p[-2]
 if y>y0:tri=[(x,y),(x-3*S,y-5*S),(x+3*S,y-5*S)]
 elif x>x0:tri=[(x,y),(x-5*S,y-3*S),(x-5*S,y+3*S)]
 else:tri=[(x,y),(x+5*S,y-3*S),(x+5*S,y+3*S)]
 d.polygon(tri,fill='black')
nodes=[(18,['Start'],True),
 (30,['Load prepared PANDORA data','Comments and normalized OCEAN scores'],False),
 (30,['Separate author partitions','50 training | 10 validation | 10 test'],False),
 (30,['Baseline comment selection','Retain all available comments per author'],False),
 (30,['BERT encoding','Maximum 256 tokens/comment; 768-dimensional vectors'],False),
 (30,['Mean pooling','One embedding vector per author'],False),
 (30,['Feature standardization','Fit on training authors; reuse for validation and test'],False),
 (30,['Train five ElasticNet regressors','One per OCEAN trait; alpha = 0.001; l1_ratio = 0.5'],False),
 (30,['Predict validation OCEAN scores','Evaluate continuous prediction performance'],False),
 (42,['Select one decision threshold per trait','Use validation predictions only','Maximize harmonic mean of F1 and specificity'],False),
 (30,['Freeze models, scaler and thresholds','Apply to test authors'],False),
 (30,['Predict continuous test OCEAN scores'],False)]
y=1;last=None
for height,lines,rounding in nodes:
 width=95 if rounding else 426
 cur=box(259,y,width,height,lines,rounding)
 if last:arrow([(259,last[1]+last[3]),(259,y)])
 last=cur;y+=height+8
# Paired continuous and binary evaluation branches.
branch_y=y+2
arrow([(259,last[1]+last[3]),(259,branch_y-5),(128,branch_y-5),(128,branch_y)])
arrow([(259,last[1]+last[3]),(259,branch_y-5),(389,branch_y-5),(389,branch_y)])
box(128,branch_y,240,37,['Continuous evaluation','MAE | RMSE | R-squared | Pearson r'])
box(389,branch_y,240,37,['Apply frozen trait thresholds','Convert scores to Low or High'])
by=branch_y+45
arrow([(389,branch_y+37),(389,by)])
box(389,by,240,39,['Binary evaluation','Accuracy | Precision | Recall','F1 | Specificity'])
save_y=by+48
arrow([(128,branch_y+37),(128,save_y-5),(259,save_y-5),(259,save_y)])
arrow([(389,by+39),(389,save_y-5),(259,save_y-5),(259,save_y)])
box(259,save_y,360,25,['Save models, thresholds and results'])
end_y=save_y+33
arrow([(259,save_y+25),(259,end_y)])
box(259,end_y,95,18,['End'],True)
im=im.crop((0,0,W*S,(end_y+19)*S));im.save(ROOT/'e1_flowchart.png',dpi=(288,288))
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.4);sec.bottom_margin=Inches(.4);sec.left_margin=Inches(.5);sec.right_margin=Inches(.5)
for sn in ['Normal','Title','Caption']:
 st=doc.styles[sn];st.font.name='Times New Roman';st.font.size=Pt(10);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(3)
doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
p=doc.add_paragraph('E1 Experiment Pipeline','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Baseline selection, BERT embeddings and ElasticNet regression');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=p.add_run();run.add_picture(str(ROOT/'e1_flowchart.png'),width=Inches(7.2))
inline=run._r.xpath('.//wp:docPr')[0];inline.set('descr','E1 flowchart from prepared PANDORA comments through BERT, mean pooling, training-only feature scaling, five ElasticNet regressors, validation threshold selection and separate continuous and binary test evaluation.')
p=doc.add_paragraph('Figure: E1 experimental pipeline. E1 uses no Q-learning or GAN augmentation.','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Scope: This figure shows the experimental workflow. The separate saved-prediction interface currently embeds only the first ten comments per author.');p.paragraph_format.space_after=Pt(0)
doc.core_properties.title='E1 Experiment Pipeline';doc.core_properties.author=''
out=ROOT.parents[1]/'output/documents/e1_experiment_flowchart.docx';doc.save(out)
print(out);print('Diagram height inches',im.height/im.width*7.2)
